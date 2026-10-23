"""Пакет ИД в формате Word (.docx): разбиение на «страницы» для того же разбора, что и PDF.

Подрядчики сдают ИД и в PDF, и в исходном Word. У Word нет страниц в файле, но формы ИД
свёрстаны так, что каждый документ пакета (акт, паспорт, протокол, схема) — отдельный
раздел или начинается с разрыва страницы. Поэтому «страница» здесь — фрагмент между
разрывами раздела/страницы. У каждой страницы:
  * text — абзацы, ячейки таблиц и надписи (текстовые поля схем) по порядку;
  * extract_tables() — таблицы как списки строк из текстов ячеек (как у pdfplumber);
  * images — встроенные картинки (сканы свидетельств о поверке) для OCR.
Объединённые ячейки Word повторяют текст в каждой клетке сетки — в отличие от PDF,
где они пустые; извлечение данных к этому устойчиво.
"""
import io
import re
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table

W_P, W_TBL, W_T, W_BR, W_TAB = qn("w:p"), qn("w:tbl"), qn("w:t"), qn("w:br"), qn("w:tab")
W_SECTPR, W_PPR = qn("w:sectPr"), qn("w:pPr")
MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
A_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
V_IMAGEDATA = "{urn:schemas-microsoft-com:vml}imagedata"
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


class DocxPage:
    def __init__(self, number: int):
        self.number = number
        self.lines: list[str] = []
        self.tables: list[list[list[str]]] = []
        self.image_ids: list[str] = []

    @property
    def text(self) -> str:
        return "\n".join(x for x in self.lines if x.strip())

    def extract_tables(self):
        return self.tables


def _walk_text(el) -> str:
    """Текст элемента без дублей: у надписей Word хранит копию для старых версий (mc:Fallback)."""
    out = []

    def rec(node):
        for ch in node:
            if ch.tag == MC_FALLBACK:
                continue
            if ch.tag == W_T:
                out.append(ch.text or "")
            elif ch.tag == W_TAB:
                out.append("\t")
            elif ch.tag == W_BR:
                out.append("\n")
            else:
                rec(ch)
                if ch.tag == W_P:          # абзацы внутри надписей — с новой строки
                    out.append("\n")
    rec(el)
    return "".join(out)


def _images(el) -> list[str]:
    ids = []
    for node in el.iter(A_BLIP):
        if node.get(R_EMBED):
            ids.append(node.get(R_EMBED))
    for node in el.iter(V_IMAGEDATA):
        if node.get(R_ID):
            ids.append(node.get(R_ID))
    return ids


def _next_section_new_page(p) -> bool:
    """Начинает ли новую страницу раздел, следующий за этим абзацем.

    Тип начала раздела хранится в sectPr следующего раздела (у последнего — в body/sectPr).
    """
    node = p.getnext()
    while node is not None:
        sp = None
        if node.tag == W_P:
            ppr = node.find(W_PPR)
            sp = ppr.find(W_SECTPR) if ppr is not None else None
        elif node.tag == W_SECTPR:
            sp = node
        if sp is not None:
            t = sp.find(qn("w:type"))
            return t is None or t.get(qn("w:val")) != "continuous"
        node = node.getnext()
    return True


def _breaks_page(p) -> tuple[bool, bool]:
    """(разрыв страницы перед текстом абзаца, конец раздела после абзаца)."""
    # Только разрывы в прогонах самого абзаца: внутри надписей схем разрывы ничего не значат
    before = any(br.get(qn("w:type")) == "page" for r in p.findall(qn("w:r")) for br in r.findall(W_BR))
    ppr = p.find(W_PPR)
    sect = ppr.find(W_SECTPR) if ppr is not None else None
    # Раздел «на текущей странице» (continuous) страницу не начинает — следующий раздел продолжает лист
    section_end = sect is not None and _next_section_new_page(p)
    if ppr is not None and ppr.find(qn("w:pageBreakBefore")) is not None:
        before = True
    return before, section_end


def read_pages(path_or_bytes) -> tuple[list[DocxPage], dict[str, bytes]]:
    """Страницы документа и картинки (по идентификатору связи)."""
    src = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes | bytearray) else path_or_bytes
    doc = Document(src)
    images = {}
    for rid, rel in doc.part.rels.items():
        if "image" in rel.reltype and not rel.is_external:
            images[rid] = rel.target_part.blob
    pages = [DocxPage(1)]

    def new_page():
        if pages[-1].lines or pages[-1].tables or pages[-1].image_ids:
            pages.append(DocxPage(len(pages) + 1))

    for ch in doc.element.body.iterchildren():
        if ch.tag == W_P:
            before, section_end = _breaks_page(ch)
            if before:
                new_page()
            text = _walk_text(ch).strip("\n")
            ppr = ch.find(W_PPR)
            if ppr is not None and ppr.find(qn("w:numPr")) is not None and text.strip():
                text = "- " + text.strip()     # маркер списка Word в тексте не хранится — как в PDF
            pages[-1].lines.append(text)
            pages[-1].image_ids += _images(ch)
            if section_end:
                new_page()
        elif ch.tag == W_TBL:
            t = Table(ch, doc)
            rows = []
            for r in t.rows:
                try:
                    rows.append([c.text.strip() or None for c in r.cells])
                except (IndexError, ValueError):   # повреждённая сетка таблицы — пропускаем строку
                    continue
            pages[-1].tables.append(rows)
            pages[-1].lines.append("\n".join(" ".join(c for c in row if c) for row in rows))
            pages[-1].image_ids += _images(ch)
    if not (pages[-1].lines or pages[-1].tables or pages[-1].image_ids):
        pages.pop()
    for i, p in enumerate(pages, 1):
        p.number = i
    return pages, images


def is_docx(data: bytes) -> bool:
    if not data.startswith(b"PK"):
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return "word/document.xml" in z.namelist()
    except zipfile.BadZipFile:
        return False


def clean(text: str) -> str:
    return re.sub(r"[ \t]+", " ", text)
